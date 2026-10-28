from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import os, io, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "cambia-esta-clave-en-produccion")
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(BASE_DIR, "instance", "prestamos.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
db = SQLAlchemy(app)

class Cliente(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    documento = db.Column(db.String(20), unique=True, nullable=False)
    direccion = db.Column(db.String(200), default="")
    email = db.Column(db.String(120), default="")
    telefono = db.Column(db.String(30), default="")
    solicitudes = db.relationship("Solicitud", backref="cliente", lazy=True)

class Solicitud(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    monto = db.Column(db.Numeric(12,2), nullable=False)
    plazo_meses = db.Column(db.Integer, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.now, nullable=False)
    estado = db.Column(db.String(30), default="Pendiente", nullable=False)
    documentos = db.Column(db.Text, default="")
    cliente_id = db.Column(db.Integer, db.ForeignKey("cliente.id"), nullable=False)
    evaluacion = db.relationship("EvaluacionRiesgo", backref="solicitud", uselist=False, cascade="all, delete-orphan")
    contrato = db.relationship("Contrato", backref="solicitud", uselist=False, cascade="all, delete-orphan")
    desembolso = db.relationship("Desembolso", backref="solicitud", uselist=False, cascade="all, delete-orphan")

class EvaluacionRiesgo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ingreso_mensual = db.Column(db.Numeric(12,2), nullable=False)
    deudas = db.Column(db.Numeric(12,2), nullable=False)
    puntuacion = db.Column(db.Integer, nullable=False)
    historial = db.Column(db.String(30), nullable=False)
    recomendacion = db.Column(db.String(20), nullable=False)
    justificacion = db.Column(db.Text, default="")
    solicitud_id = db.Column(db.Integer, db.ForeignKey("solicitud.id"), unique=True, nullable=False)

class Revision(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    estado_documentacion = db.Column(db.String(20), nullable=False)
    comentarios = db.Column(db.Text, default="")
    fecha = db.Column(db.DateTime, default=datetime.now)
    solicitud_id = db.Column(db.Integer, db.ForeignKey("solicitud.id"), nullable=False)

class Contrato(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    fecha_firma = db.Column(db.DateTime, nullable=True)
    firma_cliente = db.Column(db.Boolean, default=False)
    firma_cooperativa = db.Column(db.Boolean, default=False)
    acepto_terminos = db.Column(db.Boolean, default=False)
    solicitud_id = db.Column(db.Integer, db.ForeignKey("solicitud.id"), unique=True, nullable=False)

class Desembolso(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    monto = db.Column(db.Numeric(12,2), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.now)
    cuenta_destino = db.Column(db.String(40), nullable=False)
    solicitud_id = db.Column(db.Integer, db.ForeignKey("solicitud.id"), unique=True, nullable=False)

def money(value):
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n <= 0:
            raise ValueError
        return n
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Ingresa un monto mayor que cero.")

@app.template_filter("soles")
def soles(value):
    return f"S/ {Decimal(value):,.2f}"

@app.route("/")
def index():
    solicitudes = Solicitud.query.order_by(Solicitud.fecha.desc()).all()
    total = len(solicitudes)
    return render_template("index.html", solicitudes=solicitudes, total=total,
        pendientes=sum(s.estado == "Pendiente" for s in solicitudes),
        aprobadas=sum(s.estado == "Aprobado" for s in solicitudes),
        desembolsadas=sum(s.estado == "Desembolsado" for s in solicitudes))

@app.route("/solicitudes/nueva", methods=["GET", "POST"])
def nueva_solicitud():
    if request.method == "POST":
        try:
            nombre = request.form.get("nombre","").strip()
            documento = request.form.get("documento","").strip()
            monto = money(request.form.get("monto"))
            plazo = int(request.form.get("plazo","0"))
            ingresos = money(request.form.get("ingresos"))
            if not nombre or not documento:
                raise ValueError("Completa el nombre y documento.")
            if plazo < 1 or plazo > 120:
                raise ValueError("El plazo debe estar entre 1 y 120 meses.")
            if Cliente.query.filter_by(documento=documento).first():
                cliente = Cliente.query.filter_by(documento=documento).first()
                cliente.nombre = nombre
                cliente.direccion = request.form.get("direccion","").strip()
                cliente.email = request.form.get("email","").strip()
                cliente.telefono = request.form.get("telefono","").strip()
            else:
                cliente = Cliente(nombre=nombre, documento=documento,
                    direccion=request.form.get("direccion","").strip(),
                    email=request.form.get("email","").strip(),
                    telefono=request.form.get("telefono","").strip())
                db.session.add(cliente)
                db.session.flush()
            files = request.files.getlist("documentos")
            saved = []
            upload_dir = os.path.join(BASE_DIR, "uploads")
            os.makedirs(upload_dir, exist_ok=True)
            for f in files:
                if f and f.filename:
                    safe_name = os.path.basename(f.filename).replace("..","")
                    stamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
                    filename = stamp + "_" + safe_name
                    f.save(os.path.join(upload_dir, filename))
                    saved.append(filename)
            solicitud = Solicitud(monto=monto, plazo_meses=plazo, cliente=cliente,
                                  documentos=";".join(saved))
            # Registrar la solicitud y, si se completaron, los datos de evaluación/revisión del formulario.
            db.session.add(solicitud)
            db.session.flush()
            deudas_txt = request.form.get("deudas", "").strip()
            puntuacion_txt = request.form.get("puntuacion", "").strip()
            historial = request.form.get("historial", "").strip()
            recomendacion = request.form.get("recomendacion", "").strip()
            justificacion = request.form.get("justificacion", "").strip()
            # La evaluación se crea solo si el usuario completó todos sus campos principales.
            if any([deudas_txt, puntuacion_txt, historial, recomendacion]):
                if not all([deudas_txt, puntuacion_txt, historial, recomendacion]):
                    raise ValueError("Para guardar la evaluación inicial, completa deudas, puntuación, historial y recomendación.")
                deudas = Decimal(deudas_txt)
                puntuacion = int(puntuacion_txt)
                if deudas < 0 or puntuacion < 0 or puntuacion > 1000:
                    raise ValueError("Las deudas no pueden ser negativas y el puntaje debe estar entre 0 y 1000.")
                if historial not in ("Bueno", "Regular", "Malo") or recomendacion not in ("Aprobar", "Rechazar"):
                    raise ValueError("Selecciona un historial y una recomendación válidos.")
                db.session.add(EvaluacionRiesgo(
                    ingreso_mensual=ingresos, deudas=deudas, puntuacion=puntuacion,
                    historial=historial, recomendacion=recomendacion,
                    justificacion=justificacion, solicitud=solicitud
                ))
            estado_docs = request.form.get("documentacion", "").strip()
            comentarios_revision = request.form.get("comentarios_revision", "").strip()
            if estado_docs:
                if estado_docs not in ("Completa", "Incompleta"):
                    raise ValueError("Selecciona un estado de documentación válido.")
                db.session.add(Revision(estado_documentacion=estado_docs,
                    comentarios=comentarios_revision, solicitud_id=solicitud.id))
            db.session.commit()
            flash(f"Solicitud #{solicitud.id} registrada correctamente.", "success")
            return redirect(url_for("detalle", solicitud_id=solicitud.id))
        except (ValueError, TypeError) as e:
            db.session.rollback()
            flash(str(e) or "Revisa los datos ingresados.", "error")
    return render_template("nueva.html")

@app.route("/solicitud/<int:solicitud_id>")
def detalle(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    revisiones = Revision.query.filter_by(solicitud_id=s.id).order_by(Revision.fecha.desc()).all()
    return render_template("detalle.html", s=s, revisiones=revisiones)

@app.route("/solicitud/<int:solicitud_id>/revision", methods=["GET","POST"])
def revision(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    if request.method == "POST":
        estado = request.form.get("documentacion")
        if estado not in ("Completa","Incompleta"):
            flash("Selecciona el estado de la documentación.", "error")
        else:
            rev = Revision(estado_documentacion=estado,
                comentarios=request.form.get("comentarios","").strip(), solicitud_id=s.id)
            db.session.add(rev)
            if estado == "Incompleta":
                s.estado = "Pendiente"
            db.session.commit()
            flash("Revisión guardada.", "success")
            return redirect(url_for("detalle", solicitud_id=s.id))
    return render_template("revision.html", s=s)

@app.route("/solicitud/<int:solicitud_id>/riesgo", methods=["GET","POST"])
def riesgo(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    ev = s.evaluacion
    if request.method == "POST":
        try:
            ingresos = money(request.form.get("ingresos"))
            deudas_raw = Decimal(request.form.get("deudas","0"))
            puntuacion = int(request.form.get("puntuacion","0"))
            historial = request.form.get("historial","")
            if deudas_raw < 0 or puntuacion < 0 or puntuacion > 1000:
                raise ValueError("Las deudas no pueden ser negativas y el puntaje debe estar entre 0 y 1000.")
            if historial not in ("Bueno","Regular","Malo"):
                raise ValueError("Selecciona un historial crediticio válido.")
            recomendacion = request.form.get("recomendacion","")
            if recomendacion not in ("Aprobar","Rechazar"):
                raise ValueError("Selecciona una recomendación.")
            justificacion = request.form.get("justificacion","").strip()
            if ev:
                ev.ingreso_mensual, ev.deudas, ev.puntuacion = ingresos, deudas_raw, puntuacion
                ev.historial, ev.recomendacion, ev.justificacion = historial, recomendacion, justificacion
            else:
                ev = EvaluacionRiesgo(ingreso_mensual=ingresos, deudas=deudas_raw, puntuacion=puntuacion,
                    historial=historial, recomendacion=recomendacion, justificacion=justificacion, solicitud=s)
                db.session.add(ev)
            db.session.commit()
            flash("Evaluación de riesgo guardada. La decisión sigue siendo responsabilidad del analista autorizado.", "success")
            return redirect(url_for("detalle", solicitud_id=s.id))
        except (ValueError, TypeError, InvalidOperation) as e:
            flash(str(e) or "Verifica los datos.", "error")
    return render_template("riesgo.html", s=s, ev=ev)

@app.route("/solicitud/<int:solicitud_id>/decision", methods=["POST"])
def decision(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    accion = request.form.get("accion")
    if accion == "aprobar":
        if not s.evaluacion or s.evaluacion.recomendacion != "Aprobar":
            flash("Registra una evaluación con recomendación Aprobar antes de aprobar.", "error")
        elif not Revision.query.filter_by(solicitud_id=s.id, estado_documentacion="Completa").first():
            flash("Primero registra una revisión con documentación completa.", "error")
        else:
            s.estado = "Aprobado"
            if not s.contrato:
                db.session.add(Contrato(solicitud=s))
            db.session.commit()
            flash("Solicitud aprobada. Ya puedes gestionar el contrato.", "success")
    elif accion == "rechazar":
        s.estado = "Rechazado"
        db.session.commit()
        flash("Solicitud rechazada y registrada.", "success")
    return redirect(url_for("detalle", solicitud_id=s.id))

@app.route("/solicitud/<int:solicitud_id>/contrato", methods=["GET","POST"])
def contrato(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    if s.estado != "Aprobado":
        flash("El contrato solo está disponible para solicitudes aprobadas.", "error")
        return redirect(url_for("detalle", solicitud_id=s.id))
    c = s.contrato
    if not c:
        c = Contrato(solicitud=s)
        db.session.add(c)
        db.session.commit()
    if request.method == "POST":
        if request.form.get("acepto") != "si":
            flash("Debes aceptar los términos para registrar la aceptación.", "error")
        else:
            c.acepto_terminos = True
            c.firma_cliente = request.form.get("firma_cliente","").strip() != ""
            c.firma_cooperativa = request.form.get("firma_cooperativa","").strip() != ""
            if c.firma_cliente and c.firma_cooperativa:
                c.fecha_firma = datetime.now()
                db.session.commit()
                flash("Firmas registradas como demostración académica. No constituyen firma digital legal.", "success")
                return redirect(url_for("detalle", solicitud_id=s.id))
            db.session.commit()
            flash("Falta registrar el nombre de una de las partes.", "error")
    return render_template("contrato.html", s=s, c=c)

@app.route("/solicitud/<int:solicitud_id>/desembolso", methods=["GET","POST"])
def desembolso(solicitud_id):
    s = Solicitud.query.get_or_404(solicitud_id)
    if s.estado != "Aprobado" or not s.contrato or not s.contrato.fecha_firma:
        flash("Para desembolsar, la solicitud debe estar aprobada y el contrato debe estar firmado.", "error")
        return redirect(url_for("detalle", solicitud_id=s.id))
    if s.desembolso:
        return render_template("desembolso.html", s=s, d=s.desembolso)
    if request.method == "POST":
        cuenta = request.form.get("cuenta","").strip()
        if not cuenta:
            flash("Ingresa la cuenta de destino.", "error")
        else:
            d = Desembolso(monto=s.monto, cuenta_destino=cuenta, solicitud=s)
            s.estado = "Desembolsado"
            db.session.add(d)
            db.session.commit()
            flash("Desembolso registrado en la base de datos (simulación; no transfiere dinero real).", "success")
            return redirect(url_for("detalle", solicitud_id=s.id))
    return render_template("desembolso.html", s=s, d=None)


@app.route("/exportar.xlsx")
def exportar():
    """Genera un reporte Excel con formato profesional y resumen financiero."""
    solicitudes = Solicitud.query.order_by(Solicitud.id.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Solicitudes"

    navy = "17365D"
    blue = "1F4E78"
    light_blue = "D9EAF7"
    pale = "F3F6FA"
    white = "FFFFFF"
    green = "E2F0D9"
    orange = "FCE4D6"
    gray = "667085"
    thin_gray = Side(style="thin", color="D7DEE8")

    headers = [
        "ID Solicitud", "Cliente", "Documento", "Correo electrónico", "Teléfono",
        "Dirección", "Monto solicitado (S/)", "Plazo (meses)", "Fecha de solicitud",
        "Estado", "Ingreso mensual (S/)", "Deudas totales (S/)", "Puntaje crediticio",
        "Historial crediticio", "Recomendación", "Justificación del analista",
        "Estado de documentación", "Fecha de revisión", "Monto desembolsado (S/)",
        "Fecha de desembolso", "Cuenta de destino", "Fecha de firma de contrato"
    ]

    # Encabezado tipo informe
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = ws.cell(1, 1, "COOPERATIVA · REPORTE DE SOLICITUDES DE PRÉSTAMO")
    title.font = Font(name="Aptos Display", size=16, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle = ws.cell(2, 1, f"Reporte generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  Registros: {len(solicitudes)}")
    subtitle.font = Font(name="Aptos", size=10, italic=True, color=gray)
    subtitle.fill = PatternFill("solid", fgColor="EAF0F6")
    subtitle.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22

    header_row = 4
    for col, value in enumerate(headers, 1):
        cell = ws.cell(header_row, col, value)
        cell.font = Font(name="Aptos", bold=True, color=white, size=10)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=navy))
    ws.row_dimensions[header_row].height = 42

    for row_num, s in enumerate(solicitudes, header_row + 1):
        e, d, c = s.evaluacion, s.desembolso, s.contrato
        revisions = Revision.query.filter_by(solicitud_id=s.id).order_by(Revision.fecha.desc()).all()
        latest_revision = revisions[0] if revisions else None
        values = [
            s.id, s.cliente.nombre, s.cliente.documento, s.cliente.email, s.cliente.telefono,
            s.cliente.direccion, float(s.monto or 0), s.plazo_meses,
            s.fecha.strftime("%d/%m/%Y %H:%M") if s.fecha else "",
            s.estado, float(e.ingreso_mensual) if e else None,
            float(e.deudas) if e else None, e.puntuacion if e else None,
            e.historial if e else "", e.recomendacion if e else "",
            e.justificacion if e else "",
            latest_revision.estado_documentacion if latest_revision else "",
            latest_revision.fecha.strftime("%d/%m/%Y %H:%M") if latest_revision and latest_revision.fecha else "",
            float(d.monto) if d else None,
            d.fecha.strftime("%d/%m/%Y %H:%M") if d and d.fecha else "",
            d.cuenta_destino if d else "",
            c.fecha_firma.strftime("%d/%m/%Y %H:%M") if c and c.fecha_firma else ""
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_num, col, value)
            cell.font = Font(name="Aptos", size=10, color="263238")
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (2, 6, 16)))
            cell.border = Border(bottom=thin_gray)
            if row_num % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=pale)
        for col in (7, 11, 12, 19):
            ws.cell(row_num, col).number_format = '"S/ " #,##0.00'
        ws.cell(row_num, 13).number_format = '0'
        estado_cell = ws.cell(row_num, 10)
        if s.estado in ("Aprobado", "Desembolsado"):
            estado_cell.fill = PatternFill("solid", fgColor=green)
            estado_cell.font = Font(name="Aptos", bold=True, color="375623")
        elif s.estado == "Rechazado":
            estado_cell.fill = PatternFill("solid", fgColor=orange)
            estado_cell.font = Font(name="Aptos", bold=True, color="9C0006")

    last_row = header_row + len(solicitudes)
    if solicitudes:
        tab = Table(displayName="TablaSolicitudes", ref=f"A{header_row}:V{last_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                            showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:V{max(last_row, header_row)}"
    ws.sheet_view.showGridLines = False

    widths = {
        1: 12, 2: 27, 3: 16, 4: 28, 5: 16, 6: 30, 7: 19, 8: 15, 9: 21, 10: 17,
        11: 20, 12: 19, 13: 17, 14: 20, 15: 17, 16: 48, 17: 23, 18: 21, 19: 21,
        20: 21, 21: 23, 22: 24
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(header_row + 1, last_row + 1):
        ws.row_dimensions[row].height = 30

    # Hoja de resumen ejecutivo
    resumen = wb.create_sheet("Resumen")
    resumen.sheet_view.showGridLines = False
    resumen.merge_cells("A1:D1")
    resumen["A1"] = "RESUMEN EJECUTIVO · PRÉSTAMOS"
    resumen["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=white)
    resumen["A1"].fill = PatternFill("solid", fgColor=navy)
    resumen["A1"].alignment = Alignment(vertical="center")
    resumen.row_dimensions[1].height = 32
    resumen["A3"], resumen["B3"] = "Indicador", "Resultado"
    for cell in resumen[3][:2]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center")
    total_solicitado = sum(float(s.monto or 0) for s in solicitudes)
    total_desembolsado = sum(float(s.desembolso.monto or 0) for s in solicitudes if s.desembolso)
    metrics = [
        ("Solicitudes registradas", len(solicitudes)),
        ("Solicitudes pendientes", sum(s.estado == "Pendiente" for s in solicitudes)),
        ("Solicitudes aprobadas", sum(s.estado == "Aprobado" for s in solicitudes)),
        ("Solicitudes rechazadas", sum(s.estado == "Rechazado" for s in solicitudes)),
        ("Préstamos desembolsados", sum(s.estado == "Desembolsado" for s in solicitudes)),
        ("Total solicitado (S/)", total_solicitado),
        ("Total desembolsado (S/)", total_desembolsado),
        ("Monto promedio solicitado (S/)", total_solicitado / len(solicitudes) if solicitudes else 0),
    ]
    for r, (label, value) in enumerate(metrics, 4):
        resumen.cell(r, 1, label)
        resumen.cell(r, 2, value)
        for cell in resumen[r][:2]:
            cell.border = Border(bottom=thin_gray)
            cell.fill = PatternFill("solid", fgColor=white if r % 2 else pale)
            cell.font = Font(name="Aptos", size=11, color="263238")
        if "(S/)" in label:
            resumen.cell(r, 2).number_format = '"S/ " #,##0.00'
    resumen.column_dimensions["A"].width = 36
    resumen.column_dimensions["B"].width = 24
    resumen["A14"] = "Nota:"
    resumen["A14"].font = Font(bold=True, color=blue)
    resumen["B14"] = "Reporte administrativo. Verifica los datos antes de tomar decisiones financieras."
    resumen["B14"].alignment = Alignment(wrap_text=True, vertical="top")
    resumen.merge_cells("B14:D15")
    resumen.column_dimensions["C"].width = 18
    resumen.column_dimensions["D"].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True,
        download_name=f"reporte_prestamos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

with app.app_context():
    os.makedirs(os.path.join(BASE_DIR, "instance"), exist_ok=True)
    db.create_all()

if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=int(os.environ.get("PORT", 5000)))
